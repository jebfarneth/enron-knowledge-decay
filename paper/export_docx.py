"""Export the built paper (paper/build/paper.html) to an editable Word file.

The Word copy is for reading and editing prose; the paper itself is still
built from paper.html.tmpl, where every number comes from the results files.
Edits made in Word (with Track Changes on) are carried back into the template
by hand. Figures are the PNG renders in paper/build/png/ (made by build.py's
figure functions at 220 dpi).

Usage: uv run --no-project --with python-docx --with beautifulsoup4 python paper/export_docx.py [out.docx]
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

PAPER = Path(__file__).resolve().parent
BUILD = PAPER / "build"
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else PAPER / "Farneth_Can_Email_Reveal_Who_Outranks_Whom.docx"
FONT, MONO = "Times New Roman", "Courier New"
WIDTH = Inches(6.5)


def set_font(style, size: float, bold: bool | None = None, italic: bool | None = None) -> None:
    style.font.name = FONT
    style.font.size = Pt(size)
    fonts = style.element.rPr.rFonts
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        fonts.attrib.pop(qn(attr), None)   # built-in heading styles otherwise keep the theme's font
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attr), FONT)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    style.font.color.rgb = RGBColor(0, 0, 0)


def runs(paragraph, node, bold=False, italic=False, mono=False, small=False) -> None:
    """Add a node's text to a paragraph, keeping bold, italic and code formatting."""
    for child in node.children:
        if isinstance(child, NavigableString):
            text = re.sub(r"\s+", " ", str(child))
            if not text:
                continue
            run = paragraph.add_run(text)
            run.bold, run.italic = bold or None, italic or None
            if mono:
                run.font.name = MONO
                run.font.size = Pt(9.5)
            if small:
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
            continue
        if not isinstance(child, Tag):
            continue
        name, classes = child.name, child.get("class") or []
        if name == "br":
            paragraph.add_run().add_break()
        else:
            runs(paragraph, child, bold=bold or name in ("b", "strong"), italic=italic or name in ("i", "em"),
                 mono=mono or name == "code" or "mono" in classes, small=small or "ci" in classes)


def tidy(paragraph) -> None:
    """Trim the spaces HTML line breaks leave at the edges of a paragraph."""
    if paragraph.runs:
        paragraph.runs[0].text = paragraph.runs[0].text.lstrip()
        paragraph.runs[-1].text = paragraph.runs[-1].text.rstrip()


def para(doc, node, style=None, align=None, **fmt):
    p = doc.add_paragraph(style=style)
    runs(p, node, **fmt)
    tidy(p)
    if align is not None:
        p.alignment = align
    return p


def caption(doc, node) -> None:
    p = para(doc, node, style="Caption")
    p.paragraph_format.space_after = Pt(10)


def table(doc, node) -> None:
    rows = node.find_all("tr")
    width = max(len(r.find_all(["th", "td"])) for r in rows)
    t = doc.add_table(rows=0, cols=width)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r in rows:
        cells = t.add_row().cells
        for i, cell in enumerate(r.find_all(["th", "td"])):
            p = cells[i].paragraphs[0]
            runs(p, cell, bold=cell.name == "th")
            tidy(p)
            for run in p.runs:
                run.font.size = Pt(9)


def figure(doc, node) -> None:
    img = node.find("img")
    png = BUILD / "png" / (Path(img["src"]).stem + ".png")
    doc.add_picture(str(png), width=WIDTH if Path(img["src"]).stem != "funnel" else Inches(3.6))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = node.find("figcaption")
    if cap:
        caption(doc, cap)


def block(doc, node) -> None:
    name, classes = node.name, node.get("class") or []
    if name == "h2":
        doc.add_heading(node.get_text(" ", strip=True).replace(" ", " "), level=1)
    elif name == "h3":
        doc.add_heading(node.get_text(" ", strip=True).replace(" ", " "), level=2)
    elif name == "p":
        p = para(doc, node, small="small" in classes)
        p.paragraph_format.space_after = Pt(6)
        if "refs" in (node.parent.get("class") or []):
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.first_line_indent = Inches(-0.3)
    elif name in ("ul", "ol"):
        for li in node.find_all("li", recursive=False):
            para(doc, li, style="List Number" if name == "ol" else "List Bullet")
    elif name == "figure":
        figure(doc, node)
    elif name == "div" and "tbl" in classes:
        table(doc, node.find("table"))
        cap = node.find("div", class_="tcap")
        if cap:
            caption(doc, cap)
    elif name == "div":
        for child in node.find_all(recursive=False):
            block(doc, child)


def main() -> None:
    soup = BeautifulSoup((BUILD / "paper.html").read_text(), "html.parser")
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Inches(1))
    set_font(doc.styles["Normal"], 11)
    set_font(doc.styles["Title"], 16, bold=True)
    set_font(doc.styles["Heading 1"], 13, bold=True)
    set_font(doc.styles["Heading 2"], 11.5, bold=True)
    set_font(doc.styles["Caption"], 9, bold=False, italic=False)
    for name in ("List Bullet", "List Number"):
        set_font(doc.styles[name], 11)

    body = soup.body
    title = doc.add_paragraph(body.find("h1").get_text(" ", strip=True), style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for cls in ("author", "affil", "draft"):
        node = body.find("div", class_=cls)
        if node:
            p = para(doc, node, align=WD_ALIGN_PARAGRAPH.CENTER, bold=cls == "author")
    abstract = body.find("div", class_="abstract")
    doc.add_heading("Abstract", level=1)
    for p_node in abstract.find_all("p"):
        para(doc, p_node).paragraph_format.space_after = Pt(6)
    para(doc, body.find("p", class_="keywords")).paragraph_format.space_after = Pt(12)
    for child in body.find("div", class_="columns").find_all(recursive=False):
        block(doc, child)
    doc.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
